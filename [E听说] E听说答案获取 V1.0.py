import os as oS
import sys as Sys
import re as Reg
import datetime as Date
import json as Json
import openpyxl as Excel

def RemoveHTMLTags(Text):
    Enter1, Enter2, Clean = Reg.compile("</br>"), Reg.compile("\\s*\\|\\s*"), Reg.compile('<.*?>');
    return Reg.sub(Clean, '', Reg.sub(Enter2, '\n', Reg.sub(Enter1, '\n', Text)));


[Url, Did, FileName, NowUrl, JsonData, PNow] = ["", "", "", "", {}, 1];

print("[为人民服务] E听说答案获取器 V1.0.0");
print("E听说默认缓存根目录：%USERPROFILE%\\AppData\\Roaming\\74656D705F74656D705F74656D705F74002\\\n新版保存目录：%USERPROFILE%\\AppData\\Roaming\\ETS\\");

while True:
    Url = input("请输入缓存文件具体目录（留空则退出）：");
    Did = Url.split("\\")[-1]; # 获取 ID 名
    FileName = "[E听说答案]" + Date.datetime.now().date().isoformat() + " " + Did + ".xlsx";
    if Url == "": Sys.exit(); # 退出
    if not oS.path.exists(Url): print("未找到文件目录！"); continue; # 读取失败
    if not Reg.match(r'^\d+$', Did): print("文件目录不正确！"); continue; # 不是作业目录

    WB = Excel.Workbook(); # 创建表格
    WS = WB.active; # 获取表格内容

    [NowUrl, JsonData, PNow] = ["", {}, 1]; # 当前是第几题
    for ProblemID in oS.listdir(Url):
        NowUrl = oS.path.join(Url, ProblemID);
        if not "content" in ProblemID: continue; # 不是答案文件夹，跳过
        if not oS.path.isdir(NowUrl): continue; # 不是文件夹，跳过

        with open(NowUrl + "\\content2.json", "r", encoding="utf-8") as Str: JsonData = Json.load(Str);
        QuestionType = JsonData["structure_type"];
        AnswerList = [];

        if QuestionType == "collector.read": # 模仿朗读
            AnswerList.append([RemoveHTMLTags("标准答案：" + JsonData["info"]["value"])]);

        if QuestionType == "collector.role": # 回答问题
            for Qlist in JsonData["info"]["question"]:
                AnswerList.append([]);
                AnswerList[-1].append(RemoveHTMLTags("采分点：" + Qlist["keywords"]));
                for Anslist in Qlist["std"]: AnswerList[-1].append(RemoveHTMLTags("可选答案：" + Anslist["value"]));

        if QuestionType == "collector.3q5a": # 三问五答
            for Qlist in JsonData["info"]["question"]:
                AnswerList.append([]);
                AnswerList[-1].append(RemoveHTMLTags("采分点：" + Qlist["keywords"]));
                for Anslist in Qlist["std"]: AnswerList[-1].append(RemoveHTMLTags("可选答案：" + Anslist["value"]));

        if QuestionType == "collector.picture": # 信息转述
            AnswerList.append([]);
            if "keypoint" in JsonData["info"]: AnswerList[-1].append(RemoveHTMLTags("采分点：\n" + JsonData["info"]["keypoint"]));
            if "analyze" in JsonData["info"]: AnswerList[-1].append(RemoveHTMLTags("分析：\n" + JsonData["info"]["analyze"]));
            for Anslist in JsonData["info"]["std"]: AnswerList[-1].append(RemoveHTMLTags("可选答案：" + Anslist["value"]));

        SubP, WS.title = 1, "答案";
        WS.append(["第" + str(PNow) + "大题，编号：" + str(ProblemID) + "，题目类型" + QuestionType]);
        WS.merge_cells(start_row=WS.max_row, start_column=1, end_row=WS.max_row, end_column=3);

        for SubQue in AnswerList:
            WS.append(["", "第" + str(SubP) + "小题答案列表："]);
            for AList in SubQue: WS.append(["", "", AList]);
            WS.append([""]);
            SubP += 1;

        WS.append([""]); # 与下一题间隔几行
        WS.append([""]); # 与下一题间隔几行
        PNow += 1;

    # 设置样式
    WS.column_dimensions['A'].width = 10;
    WS.column_dimensions['B'].width = 25;
    WS.column_dimensions['C'].width = 100;
    for R in WS.iter_rows():
        for C in R: C.alignment = Excel.styles.Alignment(wrap_text=True)

    WB.save(FileName);
    print("答案获取成功！文件名：" + FileName);
